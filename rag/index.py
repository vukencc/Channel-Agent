"""Content-addressed embeddings and a reusable parent/child retrieval index."""
import hashlib
import fcntl
import json
import sqlite3
from dataclasses import dataclass
from collections import OrderedDict
from functools import lru_cache
from html.parser import HTMLParser
from core.log import get_logger
from pathlib import Path
from threading import RLock

import numpy as np

import config
from rag.chunking import TextSplitter
from rag.file_events import FileEvents
from rag.embedding import get_embedding_model, artifact_signature, local_model_path
from rag.lexical import BM25Index
from rag.result_cache import ResultCache
from rag.text_store import (TextRecord, TextSequence, text_store, shared_cache,
                            object_bytes, check_memory_budget)


logger = get_logger(__name__)
_document_cache = OrderedDict()
_document_lock = RLock()
_file_events = {}


class TextHTMLParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts, self.hidden = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in {'script', 'style'}:
            self.hidden += 1

    def handle_endtag(self, tag):
        if tag in {'script', 'style'}:
            self.hidden = max(0, self.hidden - 1)

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def load_document(path: Path) -> str:
    if path.suffix.lower() in {'.html', '.htm'}:
        parser = TextHTMLParser()
        parser.feed(path.read_text(encoding='utf-8'))
        return '\n'.join(parser.parts)
    if path.suffix.lower() == '.pdf':
        from pypdf import PdfReader
        return '\n'.join(page.extract_text() or '' for page in PdfReader(path).pages)
    if path.suffix.lower() == '.docx':
        from docx import Document
        return '\n'.join(paragraph.text for paragraph in Document(path).paragraphs)
    return path.read_text(encoding='utf-8')


@dataclass
class DocLoader:
    dir: Path
    confined: bool = False
    track_updates: bool = False

    def load(self) -> list[dict]:
        check_memory_budget()
        root = Path(self.dir).resolve()
        if not root.is_dir():
            raise ValueError(f'DOC_DIR is not a directory: {root}')
        documents = []
        fingerprints = []
        store = text_store(config.RAG_CACHE_DIR) if config.RAG_MEMORY_LIMIT_MB else None
        with _document_lock:
            cache = _document_cache.setdefault(root, {})
            events = _file_events.setdefault(root, None)
            if events is None:
                events = _file_events[root] = FileEvents()
            changed = events.changed()
            _document_cache.move_to_end(root)
            while len(_document_cache) > 4:
                evicted, _ = _document_cache.popitem(last=False)
                _file_events.pop(evicted).close()
            present = set()
            for path in sorted(root.rglob('*')):
                if self.confined and not path.resolve().is_relative_to(root):
                    continue
                if not path.is_file() or path.suffix.lower() not in {'.txt', '.md', '.html', '.htm', '.pdf', '.docx'}:
                    continue
                stat = path.stat()
                fingerprint = (stat.st_size, stat.st_mtime_ns, stat.st_ctime_ns, stat.st_ino)
                present.add(path)
                events.watch(path)
                if (changed is None or path in changed or path not in cache or cache[path][0] != fingerprint
                        or bool(store) != isinstance(cache[path][1], TextRecord)
                        or ('updated_at' in cache[path][1]['metadata']) != self.track_updates):
                    try:
                        if config.RAG_MEMORY_LIMIT_MB and (len(present) % 128 == 0 or stat.st_size > 1024 ** 2):
                            check_memory_budget(stat.st_size * 4)
                        content = load_document(path)
                    except ImportError:
                        logger.warning('缺少可选解析器，跳过 %s；PDF 需 pypdf，Word 需 python-docx', path)
                        continue
                    name = path.relative_to(root).as_posix()
                    doc = {'content': content, 'metadata': {'source': str(path), 'filename': name, 'doc_id': name}}
                    if self.track_updates:
                        doc['metadata']['updated_at'] = stat.st_mtime
                    if store:
                        doc = TextRecord({'metadata': doc['metadata']}, 'content', content, store)
                    cache[path] = (fingerprint, doc, hashlib.sha256(content.encode()).hexdigest())
                documents.append(cache[path][1])
                fingerprints.append((str(path), cache[path][2], stat.st_mtime_ns) if self.track_updates else
                                    (str(path), cache[path][2]))
            for deleted in cache.keys() - present:
                del cache[deleted]
        self.fingerprint = hashlib.sha256(json.dumps(fingerprints).encode()).hexdigest()
        if store:
            store.flush()
        check_memory_budget()
        return documents


def split_document(content, parent_chars, child_chars):
    if not config.RAG_MEMORY_LIMIT_MB:
        return _split_cached(content, parent_chars, child_chars)
    cache = shared_cache()
    key = ('split', hashlib.sha256(content.encode()).digest(), parent_chars, child_chars)
    value = cache.get(key)
    if value is None:
        check_memory_budget(len(content) * 8)
        value = _split_uncached(content, parent_chars, child_chars)
        cache.put(key, value, object_bytes(value) + object_bytes(key))
    return value


@lru_cache(maxsize=4096)
def _split_cached(content, parent_chars, child_chars):
    return _split_uncached(content, parent_chars, child_chars)


def _split_uncached(content, parent_chars, child_chars):
    splitter = TextSplitter(parent_chars, child_chars)
    return tuple((text, start, end, tuple(splitter.pack_children(text)))
                 for text, start, end in splitter.split_parents(content))


def cached_embeddings(texts: list[str], model, progress=None) -> np.ndarray:
    """Serialize builds per model across processes; completed matrices remain read-only."""
    if not texts:
        return np.empty((0, 0), dtype=np.float32)
    directory = config.RAG_CACHE_DIR / 'vectors'
    directory.mkdir(parents=True, exist_ok=True)
    identity = hashlib.sha256(model.identity.encode()).hexdigest()
    with (directory / f'{identity}.lock').open('w') as lock:
        fcntl.flock(lock, fcntl.LOCK_EX)
        return _cached_embeddings(texts, model, progress)


def _cached_embeddings(texts: list[str], model, progress=None) -> np.ndarray:
    """Persist each content vector; transactions allow safe interrupted-build resume."""
    if not texts:
        return np.empty((0, 0), dtype=np.float32)
    directory = config.RAG_CACHE_DIR / 'vectors'
    directory.mkdir(parents=True, exist_ok=True)
    identity = hashlib.sha256(model.identity.encode()).hexdigest()
    sequence = hashlib.sha256()
    for text in texts:
        sequence.update(hashlib.sha256(text.encode()).digest())
    matrix_path = directory / f'{identity}-{sequence.hexdigest()}.npy'
    complete_path = matrix_path.with_suffix('.complete')
    if matrix_path.exists() and complete_path.exists():
        return np.load(matrix_path, mmap_mode='r', allow_pickle=False)
    with sqlite3.connect(directory / f'{identity}.sqlite', timeout=120) as db:
        db.execute('CREATE TABLE IF NOT EXISTS vectors (key TEXT PRIMARY KEY, value BLOB NOT NULL)')
        matrix = None
        for start in range(0, len(texts), 512):
            batch = texts[start:start + 512]
            keys = [hashlib.sha256(text.encode()).hexdigest() for text in batch]
            unique = dict(zip(keys, batch))
            placeholders = ','.join('?' for _ in unique)
            found = dict(db.execute(f'SELECT key,value FROM vectors WHERE key IN ({placeholders})', list(unique)))
            missing = [key for key in unique if key not in found]
            if missing:
                values = model.embed_documents([unique[key] for key in missing])
                added = [(key, vector.astype(np.float32).tobytes()) for key, vector in zip(missing, values, strict=True)]
                db.executemany('INSERT OR REPLACE INTO vectors VALUES (?,?)', added)
                db.commit()
                found.update(added)
            rows = np.stack([np.frombuffer(found[key], dtype=np.float32) for key in keys])
            if matrix is None:
                matrix = np.lib.format.open_memmap(matrix_path, mode='w+', dtype=np.float32,
                                                   shape=(len(texts), rows.shape[1]))
            matrix[start:start + len(rows)] = rows
            if progress:
                progress(min(start + 512, len(texts)), len(texts), len(missing))
    matrix.flush()
    complete_path.touch()
    return np.load(matrix_path, mmap_mode='r', allow_pickle=False)


class RetrievalIndex:
    def __init__(self, documents: list[dict], progress=None, *, source_root: Path | None = None):
        check_memory_budget()
        store = text_store(config.RAG_CACHE_DIR) if config.RAG_MEMORY_LIMIT_MB else None
        self.parents: list[dict] = []
        self.children: list[dict] = []
        splitter = TextSplitter(config.RAG_PARENT_CHARS, config.RAG_CHILD_CHARS)
        seen_docs = set()
        for doc in documents:
            metadata = dict(doc['metadata'])
            doc_id = str(metadata.get('doc_id', metadata['filename']))
            if doc_id in seen_docs:
                raise ValueError(f'Duplicate document ID: {doc_id}')
            seen_docs.add(doc_id)
            for pi, (text, start, end, children) in enumerate(split_document(doc['content'], config.RAG_PARENT_CHARS, config.RAG_CHILD_CHARS)):
                parent_index = len(self.parents)
                parent_id = json.dumps([doc_id, pi], ensure_ascii=False, separators=(',', ':'))
                self.parents.append({'id': parent_id, 'document': text, 'metadata': {
                    **metadata, 'doc_id': doc_id, 'parent': pi, 'start': start, 'end': end,
                }})
                if store:
                    parent = self.parents[-1]
                    self.parents[-1] = TextRecord({'id': parent_id, 'metadata': parent['metadata']}, 'document', text, store)
                for ci, (child, a, b) in enumerate(children):
                    self.children.append({'content': child, 'parent_index': parent_index,
                                          'child': ci, 'start': start + a, 'end': start + b})
                    if store:
                        child_row = self.children[-1]
                        self.children[-1] = TextRecord({key: value for key, value in child_row.items() if key != 'content'},
                                                      'content', child, store)
            if store:
                check_memory_budget()
        self.by_id = {parent['id']: parent for parent in self.parents}
        self.parent_indexes = np.array([child['parent_index'] for child in self.children], dtype=np.int32)
        self.model = get_embedding_model() if self.children else None
        self.vectors = cached_embeddings(TextSequence(self.children, 'content') if store else
                                        [row['content'] for row in self.children], self.model, progress)
        if config.RAG_BM25_PERSIST or store:
            from rag.lexical_store import PersistentBM25
            namespace = hashlib.sha256(json.dumps([str((source_root or config.DOC_DIR).resolve()), config.RAG_PARENT_CHARS,
                                                   config.RAG_CHILD_CHARS, 'jieba-bm25-v1']).encode()).hexdigest()
            self.lexical = PersistentBM25(config.RAG_CACHE_DIR / 'lexical' / f'{namespace}.sqlite',
                TextSequence(self.parents, 'document'), [parent['id'] for parent in self.parents])
        else:
            self.lexical = BM25Index([parent['document'] for parent in self.parents])
        self.query_cache = ResultCache()
        if store:
            store.flush()
        check_memory_budget()

    def dense(self, query: str, limit: int, allowed: set[int] | None = None) -> list[dict]:
        check_memory_budget()
        if not self.children:
            return []
        key = hashlib.sha256(query.encode()).digest()
        vector = self.query_cache.get(key, config.RAG_QUERY_CACHE_SIZE)
        if vector is None:
            vector = self.model.embed_queries([query])[0]
            self.query_cache.put(key, vector, config.RAG_QUERY_CACHE_SIZE)
        if self.vectors.shape[1] != len(vector):
            raise ValueError('Query/document embedding dimensions differ')
        self.vector_backend = 'exact'
        if allowed is None and config.RAG_VECTOR_BACKEND == 'ann' and len(self.children) >= config.RAG_ANN_MIN_CHILDREN:
            try:
                from rag.vector_backend import AnnIndex
                if not hasattr(self, 'ann_index'):
                    with _lock:
                        if not hasattr(self, 'ann_index'):
                            self.ann_index = AnnIndex(self.vectors, config.RAG_CACHE_DIR / 'ann')
                k = min(len(self.children), max(limit * 4, config.RAG_ANN_EF_SEARCH))
                while True:
                    hits = self.ann_index.search(vector, k)
                    scores = {}
                    for child, score in hits:
                        parent = int(self.parent_indexes[child])
                        scores[parent] = max(scores.get(parent, -np.inf), score)
                    if len(scores) >= min(limit, len(self.parents)) or k == len(self.children):
                        break
                    k = min(len(self.children), k * 2)
                self.vector_backend = 'ann'
                order = sorted(scores, key=lambda i: (-scores[i], i))[:limit]
                return [{'id': self.parents[i]['id'], 'score': scores[i]} for i in order]
            except (OSError, RuntimeError, ValueError) as exc:
                logger.warning('ANN 不可用，显式回退 exact：%s', exc)
        child_scores = self.vectors @ vector
        parent_scores = np.full(len(self.parents), -np.inf, dtype=np.float32)
        np.maximum.at(parent_scores, self.parent_indexes, child_scores)
        # Parent list is deterministic; ties retain document/chunk order.
        if allowed is None:
            order = np.argsort(-parent_scores, kind='stable')[:limit]
        else:
            eligible = np.array(sorted(allowed), dtype=np.int64)
            order = eligible[np.argsort(-parent_scores[eligible], kind='stable')[:limit]]
        return [{'id': self.parents[i]['id'], 'score': float(parent_scores[i])} for i in order]

    def bm25(self, query: str, limit: int, allowed: set[int] | None = None) -> list[dict]:
        return [{'id': self.parents[i]['id'], 'score': score}
                for i, score in self.lexical.search(query, len(self.parents) if allowed is not None else limit)
                if allowed is None or i in allowed][:limit]


_lock = RLock()
_cached_key = None
_cached_index = None


def get_index(root: Path | None = None, *, track_updates: bool = False) -> RetrievalIndex:
    """Read content to detect same-size/mtime edits; reuse all expensive work."""
    global _cached_key, _cached_index
    loader = DocLoader(root or config.DOC_DIR, confined=root is not None, track_updates=track_updates)
    documents = loader.load()
    key = (str(root or config.DOC_DIR), root is not None, track_updates, loader.fingerprint, config.RAG_PARENT_CHARS,
           config.RAG_CHILD_CHARS, config.EMBEDDING_MODEL_SOURCE,
           config.EMBEDDING_MODEL_URL, config.EMBEDDING_MODEL_NAME,
           config.EMBEDDING_MODEL_API_KEY, config.EMBEDDING_LOCAL_PATH,
           config.RAG_CACHE_DIR, config.RAG_BATCH_SIZE, config.RAG_THREADS)
    key += (config.RAG_VECTOR_BACKEND, config.RAG_ANN_MIN_CHILDREN, config.RAG_ANN_M,
            config.RAG_ANN_EF_CONSTRUCTION, config.RAG_ANN_EF_SEARCH, config.RAG_BM25_PERSIST,
            config.RAG_MEMORY_LIMIT_MB)
    if config.EMBEDDING_MODEL_SOURCE.upper() == 'LOCAL':
        key += (artifact_signature(local_model_path()),)
    if key == _cached_key:
        return _cached_index
    with _lock:
        if key != _cached_key:
            index = RetrievalIndex(documents, source_root=root) if root is not None else RetrievalIndex(documents)
            _cached_key, _cached_index = key, index
        return _cached_index
